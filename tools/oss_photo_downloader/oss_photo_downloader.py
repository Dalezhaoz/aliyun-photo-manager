"""Windows desktop bulk downloader for Alibaba Cloud OSS photos.

The spreadsheet must contain an ID column.  Each object is looked up as:
    <object prefix>/<ID><extension>
and downloaded to:
    <chosen output folder>/<group-column value>/<ID><extension>
"""

from __future__ import annotations

import csv
import base64
import hashlib
import hmac
import queue
import re
import threading
from email.utils import formatdate
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit
from urllib.request import Request, urlopen

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


INVALID_PATH_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
ID_PATTERN = re.compile(r"^\d{15}(?:\d{2}[0-9Xx])?$")


@dataclass(frozen=True)
class DownloadRow:
    row_number: int
    identity: str
    group: str


@dataclass(frozen=True)
class OssClient:
    """Minimal OSS V1 signed GET client using only Python's standard library.

    This deliberately avoids the oss2/cryptography dependency, whose source build
    fails on some Windows ARM64 Python installations.
    """
    endpoint: str
    bucket: str
    access_key_id: str
    access_key_secret: str
    security_token: str = ""

    def download(self, object_key: str, target: Path) -> None:
        endpoint = self.endpoint if "://" in self.endpoint else "https://" + self.endpoint
        parsed = urlsplit(endpoint)
        if not parsed.scheme or not parsed.netloc:
            raise ValueError("Endpoint 格式不正确。示例：https://oss-cn-hangzhou.aliyuncs.com")
        date = formatdate(usegmt=True)
        oss_headers: dict[str, str] = {}
        if self.security_token:
            oss_headers["x-oss-security-token"] = self.security_token
        canonical_headers = "".join(f"{key}:{oss_headers[key]}\n" for key in sorted(oss_headers))
        canonical_resource = f"/{self.bucket}/{object_key}"
        string_to_sign = f"GET\n\n\n{date}\n{canonical_headers}{canonical_resource}"
        signature = base64.b64encode(hmac.new(self.access_key_secret.encode("utf-8"), string_to_sign.encode("utf-8"), hashlib.sha1).digest()).decode("ascii")
        base_path = parsed.path.rstrip("/")
        encoded_key = quote(object_key, safe="/-_.~")
        url = f"{parsed.scheme}://{self.bucket}.{parsed.netloc}{base_path}/{encoded_key}"
        headers = {"Date": date, "Authorization": f"OSS {self.access_key_id}:{signature}", **oss_headers}
        request = Request(url, headers=headers, method="GET")
        partial = target.with_suffix(target.suffix + ".part")
        try:
            with urlopen(request, timeout=90) as response, partial.open("wb") as output:
                while chunk := response.read(1024 * 1024):
                    output.write(chunk)
            partial.replace(target)
        except Exception:
            partial.unlink(missing_ok=True)
            raise


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    # Prevent 18-digit IDs being displayed as 123...E+17 after Excel conversion.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def safe_folder(value: str) -> str:
    value = INVALID_PATH_CHARS.sub("_", value).strip(". ")
    return value or "未分类"


def normalize_prefix(prefix: str) -> str:
    return prefix.strip().strip("/")


def read_headers_and_rows(file_path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, str]], list[str]]:
    """Read an xlsx/xlsm or csv, preserving header text and values as text."""
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        with file_path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            headers = [h.strip() for h in (reader.fieldnames or []) if h and h.strip()]
            return headers, [{h: clean_cell(row.get(h)) for h in headers} for row in reader], []

    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx、.xlsm 或 .csv 文件。")
    if load_workbook is None:
        raise RuntimeError("缺少 openpyxl。请按 README 安装依赖。")

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheets = workbook.sheetnames
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.active.title]
    rows = worksheet.iter_rows(values_only=True)
    first = next(rows, None)
    if not first:
        raise ValueError("表格为空。")
    headers = [clean_cell(value) for value in first]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("首行必须是非空且不重复的列名。")
    result = []
    for values in rows:
        result.append({headers[i]: clean_cell(values[i] if i < len(values) else None) for i in range(len(headers))})
    return headers, result, sheets


class OssPhotoDownloader(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("阿里云 OSS 照片下载工具")
        self.geometry("820x700")
        self.minsize(760, 620)
        self.records: list[dict[str, str]] = []
        self.headers: list[str] = []
        self.event_queue: queue.Queue[tuple[str, Any]] = queue.Queue()
        self.running = False

        self.file_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.id_column_var = tk.StringVar()
        self.group_column_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.endpoint_var = tk.StringVar(value="https://oss-cn-hangzhou.aliyuncs.com")
        self.bucket_var = tk.StringVar()
        self.prefix_var = tk.StringVar()
        self.extensions_var = tk.StringVar(value=".jpg,.jpeg,.png")
        self.access_key_var = tk.StringVar()
        self.secret_key_var = tk.StringVar()
        self.token_var = tk.StringVar()
        self.workers_var = tk.IntVar(value=5)
        self.status_var = tk.StringVar(value="请选择表格。")
        self._build_ui()
        self.after(100, self._drain_events)

    def _build_ui(self) -> None:
        outer = ttk.Frame(self, padding=14)
        outer.pack(fill="both", expand=True)
        outer.columnconfigure(1, weight=1)

        ttk.Label(outer, text="阿里云 OSS 照片下载", font=("Microsoft YaHei UI", 15, "bold")).grid(column=0, row=0, columnspan=3, sticky="w", pady=(0, 12))
        self._path_row(outer, 1, "Excel / CSV 表格", self.file_var, "选择表格", self.choose_file)
        ttk.Label(outer, text="工作表").grid(column=0, row=2, sticky="w", pady=5)
        self.sheet_box = ttk.Combobox(outer, textvariable=self.sheet_var, state="disabled")
        self.sheet_box.grid(column=1, row=2, sticky="ew", pady=5)
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _: self.reload_sheet())

        ttk.Separator(outer).grid(column=0, row=3, columnspan=3, sticky="ew", pady=8)
        ttk.Label(outer, text="身份证号列").grid(column=0, row=4, sticky="w", pady=5)
        self.id_box = ttk.Combobox(outer, textvariable=self.id_column_var, state="readonly")
        self.id_box.grid(column=1, row=4, sticky="ew", pady=5)
        ttk.Label(outer, text="归档文件夹列").grid(column=0, row=5, sticky="w", pady=5)
        self.group_box = ttk.Combobox(outer, textvariable=self.group_column_var, state="readonly")
        self.group_box.grid(column=1, row=5, sticky="ew", pady=5)
        self._path_row(outer, 6, "下载保存位置", self.output_var, "选择文件夹", self.choose_output)

        ttk.Separator(outer).grid(column=0, row=7, columnspan=3, sticky="ew", pady=8)
        self._entry_row(outer, 8, "OSS Endpoint", self.endpoint_var, "例如：https://oss-cn-hangzhou.aliyuncs.com")
        self._entry_row(outer, 9, "Bucket 名称", self.bucket_var, "例如：my-photo-bucket")
        self._entry_row(outer, 10, "对象前缀（可空）", self.prefix_var, "例如：photos/2026")
        self._entry_row(outer, 11, "照片扩展名", self.extensions_var, "按顺序尝试，以英文逗号分隔")
        self._entry_row(outer, 12, "AccessKey ID", self.access_key_var, "建议使用仅有读取权限的 RAM 用户")
        self._entry_row(outer, 13, "AccessKey Secret", self.secret_key_var, "本次运行仅驻留内存", show="*")
        self._entry_row(outer, 14, "Security Token（可空）", self.token_var, "使用 STS 临时凭证时填写", show="*")

        ttk.Label(outer, text="并发下载数").grid(column=0, row=15, sticky="w", pady=5)
        ttk.Spinbox(outer, from_=1, to=20, textvariable=self.workers_var, width=8).grid(column=1, row=15, sticky="w", pady=5)
        self.start_button = ttk.Button(outer, text="开始下载", command=self.start_download)
        self.start_button.grid(column=2, row=15, sticky="e", pady=5)

        self.progress = ttk.Progressbar(outer, mode="determinate")
        self.progress.grid(column=0, row=16, columnspan=3, sticky="ew", pady=(10, 4))
        ttk.Label(outer, textvariable=self.status_var).grid(column=0, row=17, columnspan=3, sticky="w")
        self.log = tk.Text(outer, height=11, wrap="word", state="disabled")
        self.log.grid(column=0, row=18, columnspan=3, sticky="nsew", pady=(8, 0))
        outer.rowconfigure(18, weight=1)

    def _entry_row(self, parent: ttk.Frame, row: int, label: str, variable: tk.StringVar, hint: str, show: str | None = None) -> None:
        ttk.Label(parent, text=label).grid(column=0, row=row, sticky="w", pady=5)
        entry = ttk.Entry(parent, textvariable=variable, show=show or "")
        entry.grid(column=1, row=row, sticky="ew", pady=5)
        ttk.Label(parent, text=hint, foreground="#666666").grid(column=2, row=row, sticky="w", padx=(10, 0))

    def _path_row(self, parent: ttk.Frame, row: int, label: str, variable: tk.StringVar, button: str, command: Any) -> None:
        ttk.Label(parent, text=label).grid(column=0, row=row, sticky="w", pady=5)
        ttk.Entry(parent, textvariable=variable, state="readonly").grid(column=1, row=row, sticky="ew", pady=5)
        ttk.Button(parent, text=button, command=command).grid(column=2, row=row, sticky="e", padx=(10, 0))

    def choose_file(self) -> None:
        path = filedialog.askopenfilename(title="选择人员表格", filetypes=[("表格文件", "*.xlsx *.xlsm *.csv"), ("所有文件", "*.*")])
        if not path:
            return
        self.file_var.set(path)
        try:
            _, _, sheets = read_headers_and_rows(Path(path))
            if sheets:
                self.sheet_box.configure(values=sheets, state="readonly")
                self.sheet_var.set(sheets[0])
            else:
                self.sheet_box.configure(values=[], state="disabled")
                self.sheet_var.set("")
            self.reload_sheet()
        except Exception as exc:
            messagebox.showerror("无法读取表格", str(exc))

    def reload_sheet(self) -> None:
        try:
            self.headers, self.records, _ = read_headers_and_rows(Path(self.file_var.get()), self.sheet_var.get() or None)
            for box in (self.id_box, self.group_box):
                box.configure(values=self.headers)
            if self.headers:
                self.id_column_var.set(self.headers[0])
                self.group_column_var.set(self.headers[1] if len(self.headers) > 1 else self.headers[0])
            self.status_var.set(f"已加载 {len(self.records)} 行，可选择对应列。")
        except Exception as exc:
            messagebox.showerror("无法读取工作表", str(exc))

    def choose_output(self) -> None:
        path = filedialog.askdirectory(title="选择照片下载保存位置")
        if path:
            self.output_var.set(path)

    def start_download(self) -> None:
        if load_workbook is None:
            messagebox.showerror("缺少依赖", "未安装 openpyxl。请按 README 安装依赖。")
            return
        required = {
            "表格": self.file_var.get(), "身份证号列": self.id_column_var.get(), "归档文件夹列": self.group_column_var.get(),
            "下载保存位置": self.output_var.get(), "Endpoint": self.endpoint_var.get(), "Bucket": self.bucket_var.get(),
            "AccessKey ID": self.access_key_var.get(), "AccessKey Secret": self.secret_key_var.get(),
        }
        missing = [key for key, value in required.items() if not value.strip()]
        if missing:
            messagebox.showwarning("信息不完整", "请填写：" + "、".join(missing))
            return
        if self.id_column_var.get() == self.group_column_var.get():
            if not messagebox.askyesno("确认", "身份证号列与文件夹列相同，仍要继续吗？"):
                return
        self.running = True
        self.start_button.configure(state="disabled")
        self.progress.configure(maximum=len(self.records), value=0)
        self._write_log("开始下载。密钥不会写入日志或结果表。")
        task = {
            "endpoint": self.endpoint_var.get().strip(),
            "bucket": self.bucket_var.get().strip(),
            "access_key_id": self.access_key_var.get().strip(),
            "access_key_secret": self.secret_key_var.get().strip(),
            "security_token": self.token_var.get().strip(),
            "id_column": self.id_column_var.get(),
            "group_column": self.group_column_var.get(),
            "records": list(self.records),
            "extensions": self.extensions_var.get(),
            "output": self.output_var.get(),
            "prefix": self.prefix_var.get(),
            "workers": self.workers_var.get(),
        }
        threading.Thread(target=self._download_all, args=(task,), daemon=True).start()

    def _download_all(self, task: dict[str, Any]) -> None:
        try:
            client = OssClient(
                endpoint=task["endpoint"], bucket=task["bucket"],
                access_key_id=task["access_key_id"], access_key_secret=task["access_key_secret"],
                security_token=task["security_token"],
            )
            rows = [
                DownloadRow(i + 2, clean_cell(row[task["id_column"]]), clean_cell(row[task["group_column"]]))
                for i, row in enumerate(task["records"])
            ]
            extensions = [
                value.strip() if value.strip().startswith(".") else "." + value.strip()
                for value in task["extensions"].split(",") if value.strip()
            ]
            if not extensions:
                raise ValueError("请至少填写一个照片扩展名。")
            output = Path(task["output"])
            output.mkdir(parents=True, exist_ok=True)
            results: list[dict[str, str]] = []
            with ThreadPoolExecutor(max_workers=max(1, min(20, task["workers"]))) as executor:
                futures = [
                    executor.submit(self._download_one, client, output, normalize_prefix(task["prefix"]), extensions, row)
                    for row in rows
                ]
                for future in as_completed(futures):
                    result = future.result()
                    results.append(result)
                    self.event_queue.put(("row", result))
            self._write_report(output, results)
            self.event_queue.put(("done", results))
        except Exception as exc:
            self.event_queue.put(("fatal", str(exc)))

    @staticmethod
    def _download_one(client: OssClient, output: Path, prefix: str, extensions: list[str], row: DownloadRow) -> dict[str, str]:
        if not row.identity:
            return {"行号": str(row.row_number), "身份证号": "", "归档文件夹": row.group, "状态": "跳过", "说明": "身份证号为空"}
        if not ID_PATTERN.fullmatch(row.identity):
            return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "跳过", "说明": "身份证号格式异常"}
        folder = output / safe_folder(row.group)
        folder.mkdir(parents=True, exist_ok=True)
        for extension in extensions:
            object_key = "/".join(part for part in (prefix, row.identity + extension) if part)
            target = folder / (row.identity + extension)
            if target.exists() and target.stat().st_size > 0:
                return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "已存在", "说明": str(target)}
            try:
                client.download(object_key, target)
                return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "成功", "说明": object_key}
            except HTTPError as exc:
                if exc.code == 404:
                    continue
                return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "失败", "说明": f"HTTP {exc.code}: {exc.reason}"}
            except URLError as exc:
                return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "失败", "说明": f"网络错误: {exc.reason}"}
            except Exception as exc:
                return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "失败", "说明": str(exc)}
        return {"行号": str(row.row_number), "身份证号": row.identity, "归档文件夹": row.group, "状态": "未找到", "说明": "已尝试：" + ", ".join(extensions)}

    @staticmethod
    def _write_report(output: Path, results: list[dict[str, str]]) -> None:
        report = output / "下载结果.csv"
        with report.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["行号", "身份证号", "归档文件夹", "状态", "说明"])
            writer.writeheader()
            writer.writerows(sorted(results, key=lambda x: int(x["行号"])))

    def _drain_events(self) -> None:
        try:
            while True:
                kind, payload = self.event_queue.get_nowait()
                if kind == "row":
                    self.progress.configure(value=self.progress["value"] + 1)
                    self.status_var.set(f"已处理 {int(self.progress['value'])}/{int(self.progress['maximum'])} 行")
                    self._write_log(f"第 {payload['行号']} 行：{payload['状态']} — {payload['身份证号']} ({payload['说明']})")
                elif kind == "done":
                    counts: dict[str, int] = {}
                    for item in payload:
                        counts[item["状态"]] = counts.get(item["状态"], 0) + 1
                    self.status_var.set("完成。" + "，".join(f"{k} {v}" for k, v in counts.items()) + "；结果已保存为 下载结果.csv")
                    self._write_log("下载完成。")
                    self.running = False
                    self.start_button.configure(state="normal")
                elif kind == "fatal":
                    self.running = False
                    self.start_button.configure(state="normal")
                    self.status_var.set("下载未完成。")
                    self._write_log("错误：" + payload)
                    messagebox.showerror("下载失败", payload)
        except queue.Empty:
            pass
        self.after(100, self._drain_events)

    def _write_log(self, message: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", message + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")


if __name__ == "__main__":
    OssPhotoDownloader().mainloop()
