from __future__ import annotations

import base64
import html
import json
import math
import re
import sys
from io import BytesIO
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path
from typing import Iterable
from urllib.parse import quote

from openpyxl import load_workbook


DOC_TYPE_SEAT = "seat"
DOC_TYPE_DOOR = "door"
DOC_TYPE_DESK = "desk"

PLACEHOLDER_PATTERN = re.compile(r"\$\{([^}]+)\}")
TEMPLATE_STORAGE_NAME = ".exam_print_templates.json"
SUPPORTED_IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
IMAGE_MIME_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
}


@dataclass
class PrintTemplate:
    name: str
    doc_type: str
    main_html: str
    item_html: str = ""
    settings: dict[str, int | str] = field(default_factory=dict)
    builtin: bool = False


@dataclass(frozen=True)
class PrintDataConfig:
    excel_path: Path
    room_column: str
    seat_column: str
    exam_no_column: str
    site_column: str
    subject_column: str
    unit_column: str
    job_column: str
    photo_dir: Path | None = None
    photo_match_column: str = ""


def template_storage_path() -> Path:
    if getattr(sys, "frozen", False):
        app_root = Path(sys.executable).resolve().parent
        if app_root.name.lower() == "app" and (app_root.parent / "aliyun_photo_manager_launcher.exe").exists():
            return app_root.parent / TEMPLATE_STORAGE_NAME
        return app_root / TEMPLATE_STORAGE_NAME
    return Path(__file__).resolve().parents[2] / TEMPLATE_STORAGE_NAME


def _escape(value: object) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def _coerce_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _to_file_url(path: Path) -> str:
    resolved = path.resolve()
    return "file:///" + quote(str(resolved).replace("\\", "/"), safe="/:")


def _to_image_data_uri(path: Path) -> str:
    try:
        from PIL import Image, ImageFile

        ImageFile.LOAD_TRUNCATED_IMAGES = True

        with Image.open(path) as image:
            if image.mode not in {"RGB", "RGBA"}:
                image = image.convert("RGB")
            image.thumbnail((360, 480))
            output = BytesIO()
            image.save(output, format="PNG")
        encoded = base64.b64encode(output.getvalue()).decode("ascii")
        return f"data:image/png;base64,{encoded}"
    except Exception:
        mime_type = IMAGE_MIME_TYPES.get(path.suffix.lower())
        if not mime_type:
            return _to_file_url(path)
    try:
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    except OSError:
        return _to_file_url(path)
    return f"data:{mime_type};base64,{encoded}"


def available_placeholders(headers: Iterable[str]) -> list[str]:
    base = sorted({str(header).strip() for header in headers if str(header).strip()})
    extras = [
        "考场",
        "考点",
        "考试科目",
        "起始考号",
        "结束考号",
        "单位岗位汇总",
        "考生人数",
        "照片",
        "照片标签",
        "面试报到时间",
        "考生列表",
        "座次表表格",
        "桌贴列表",
    ]
    return [f"${{{item}}}" for item in [*base, *extras]]


def load_excel_headers(excel_path: Path) -> list[str]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [_coerce_text(value) for value in header_row if _coerce_text(value)]
    finally:
        workbook.close()


def load_excel_records(excel_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, ())
        headers = [_coerce_text(value) for value in header_row]
        records: list[dict[str, str]] = []
        for row in rows:
            record: dict[str, str] = {}
            has_value = False
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = _coerce_text(row[index] if index < len(row) else "")
                record[header] = value
                has_value = has_value or bool(value)
            if has_value:
                records.append(record)
        return records
    finally:
        workbook.close()


def load_templates() -> dict[str, list[PrintTemplate]]:
    templates = default_templates()
    storage_path = template_storage_path()
    if not storage_path.exists():
        return templates
    try:
        payload = json.loads(storage_path.read_text(encoding="utf-8"))
    except Exception:
        return templates
    if not isinstance(payload, dict):
        return templates
    for doc_type, items in payload.items():
        if not isinstance(items, list):
            continue
        loaded: list[PrintTemplate] = []
        for item in items:
            if not isinstance(item, dict):
                continue
            name = _coerce_text(item.get("name"))
            if not name:
                continue
            loaded.append(
                PrintTemplate(
                    name=name,
                    doc_type=doc_type,
                    main_html=str(item.get("main_html", "")),
                    item_html=str(item.get("item_html", "")),
                    settings=dict(item.get("settings", {})),
                    builtin=False,
                )
            )
        if loaded:
            builtin_names = {template.name for template in templates.get(doc_type, [])}
            merged = templates.get(doc_type, [])[:]
            merged.extend(template for template in loaded if template.name not in builtin_names)
            templates[doc_type] = merged
    return templates


def save_template(template: PrintTemplate) -> None:
    storage_path = template_storage_path()
    payload: dict[str, list[dict[str, object]]] = {}
    if storage_path.exists():
        try:
            existing = json.loads(storage_path.read_text(encoding="utf-8"))
            if isinstance(existing, dict):
                payload = existing
        except Exception:
            payload = {}
    items = [item for item in payload.get(template.doc_type, []) if item.get("name") != template.name]
    items.append(
        {
            "name": template.name,
            "main_html": template.main_html,
            "item_html": template.item_html,
            "settings": template.settings,
        }
    )
    payload[template.doc_type] = items
    storage_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def default_templates() -> dict[str, list[PrintTemplate]]:
    written_top = """
<div style="text-align:center;font-size:12px;font-weight:bold;color:#444;">
考点：${考点} | 考场：${考场}
</div>
<div style="text-align:center;font-size:9px;color:#555;">
准考证号起止：${起始考号} - ${结束考号} | 共有考生：${考生人数} | 所含专业：${考试科目} | 监考人员签字：________________
</div>
""".strip()
    written_item = """
<div>姓名：${姓名}</div>
<div>考号：${考号}</div>
<div>考场：${考场}  座号：${座号}</div>
<div>进场签字：</div>
<div>出场签字：</div>
""".strip()
    interview_top = """
<div style="text-align:center;font-size:22px;font-weight:bold;color:#1f2937;">
2025年菏泽市牡丹区公开招聘教师面试考生签到表
</div>
<div style="text-align:center;font-size:18px;font-weight:bold;color:#b4535a;">
（${考点}）
</div>
""".strip()
    interview_item = """
<div>姓名：${姓名}</div>
<div>身份证号：</div>
<div>${身份证号}</div>
<div>面试序号：${考号}</div>
<div>候考室：${考场}</div>
<div>报考专业：${报考岗位}</div>
<div>进场签名：</div>
""".strip()
    seat_main = """
<html>
<head>
<style>
body { font-family: 'Microsoft YaHei UI'; margin: 10px; color: #111827; }
.candidate-grid { width: 100%; border-collapse: collapse; table-layout: fixed; }
.candidate-grid > tr > td, .candidate-grid > tbody > tr > td { border: 1px solid #444; padding: 0; height: 94px; vertical-align: top; }
.candidate-card { width: 100%; height: 94px; border-collapse: collapse; table-layout: fixed; font-size: 12px; line-height: 1.35; }
.candidate-card td { border: none; padding: 0; vertical-align: top; }
.candidate-card .photo-cell { width: 68px; }
.candidate-card .info-cell { padding: 3px 5px; }
.photo { width: 68px; height: 94px; object-fit: cover; display: block; }
.footer-title { margin: 10px 10px 0; padding: 8px 12px; text-align: center; color: #c62828; background: #f5dada; border-radius: 8px; font-size: 18px; }
</style>
</head>
<body>
  ${考生列表}
  <div class="footer-title">2026年度周村区（文昌湖区）事业单位公开招聘综合类岗位人员面试</div>
</body>
</html>
""".strip()
    seat_item = """
<table class="candidate-card">
  <tr>
    <td class="photo-cell">${照片标签}</td>
    <td class="info-cell">
      <div>姓名:${姓名}</div>
      <div>身份证号:</div>
      <div>${身份证号}</div>
      <div>报考单位:${报考单位}</div>
      <div>报考职位:${报考岗位}</div>
    </td>
  </tr>
</table>
""".strip()
    door_main = """
<html>
<head>
<style>
body { font-family: 'Microsoft YaHei UI'; margin: 32px; color: #111827; }
.door { border: 2px solid #cbd5e1; padding: 24px 28px; }
.title { text-align: center; font-size: 34px; font-weight: 700; margin-bottom: 18px; }
.line { font-size: 18px; line-height: 2; }
</style>
</head>
<body>
  <div class="door">
    <div class="title">考场门贴</div>
    <div class="line">考点：${考点}</div>
    <div class="line">考场：${考场}</div>
    <div class="line">考试科目：${考试科目}</div>
    <div class="line">准考证号：${起始考号} - ${结束考号}</div>
    <div class="line">单位岗位：${单位岗位汇总}</div>
    <div class="line">人数：${考生人数}</div>
  </div>
</body>
</html>
""".strip()
    desk_main = """
<html>
<head>
<style>
body { font-family: 'Microsoft YaHei UI'; margin: 18px; color: #111827; }
.page { page-break-after: always; }
.page:last-child { page-break-after: auto; }
.desk-grid { width: 100%; border-collapse: separate; border-spacing: 10px; table-layout: fixed; }
.desk-grid td { border: 1px dashed #94a3b8; padding: 10px; vertical-align: top; height: 100px; }
.desk-card { font-size: 12px; line-height: 1.6; }
.desk-name { font-size: 16px; font-weight: 700; margin-bottom: 6px; }
</style>
</head>
<body>
  ${桌贴列表}
</body>
</html>
""".strip()
    desk_item = """
<div class="desk-card">
  <div class="desk-name">${姓名}</div>
  <div>座号：${座号}</div>
  <div>准考证号：${考号}</div>
  <div>${报考单位}${报考岗位}</div>
</div>
""".strip()
    return {
        DOC_TYPE_SEAT: [
            PrintTemplate(
                "笔试座次表",
                DOC_TYPE_SEAT,
                seat_main,
                written_item,
                {
                    "columns": 5,
                    "layout": "candidate_grid",
                    "top_html": written_top,
                    "bottom_html": "",
                    "photo_width": 58,
                    "font_size": 6,
                    "line_spacing": 7,
                },
                builtin=True,
            ),
            PrintTemplate(
                "面试签到表",
                DOC_TYPE_SEAT,
                seat_main,
                interview_item,
                {
                    "columns": 5,
                    "layout": "candidate_grid",
                    "top_html": interview_top,
                    "bottom_html": "",
                    "photo_width": 52,
                    "font_size": 7,
                    "line_spacing": 8,
                },
                builtin=True,
            ),
        ],
        DOC_TYPE_DOOR: [
            PrintTemplate("标准门贴", DOC_TYPE_DOOR, door_main, "", {}, builtin=True),
        ],
        DOC_TYPE_DESK: [
            PrintTemplate("桌贴 10 人版", DOC_TYPE_DESK, desk_main, desk_item, {"columns": 2, "items_per_page": 10}, builtin=True),
            PrintTemplate("桌贴 30 人版", DOC_TYPE_DESK, desk_main, desk_item, {"columns": 5, "items_per_page": 30}, builtin=True),
        ],
    }


def room_values(records: Iterable[dict[str, str]], room_column: str) -> list[str]:
    values = sorted({_coerce_text(record.get(room_column, "")) for record in records if _coerce_text(record.get(room_column, ""))})
    return values


def build_room_context(
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    *,
    sort_column: str = "",
    image_mode: str = "data",
) -> tuple[dict[str, str], list[dict[str, str]]]:
    if config.room_column and room_value:
        room_records = [record for record in records if _coerce_text(record.get(config.room_column, "")) == room_value]
    else:
        room_records = list(records)
    if not room_records:
        raise ValueError("没有匹配到任何考生。")

    if sort_column:
        sorted_records = sorted(room_records, key=lambda item: _natural_sort_key(item.get(sort_column, "")))
    else:
        sorted_records = sorted(room_records, key=lambda item: _sort_key(item, config))
    for record in sorted_records:
        record["照片"] = _match_photo(record, config, image_mode=image_mode)
        record["座号"] = record.get(config.seat_column, "") if config.seat_column else record.get("座号", "")
        record["考号"] = record.get(config.exam_no_column, "") if config.exam_no_column else record.get("考号", "")
        record["报考单位"] = record.get(config.unit_column, "") if config.unit_column else record.get("报考单位", "")
        record["报考岗位"] = record.get(config.job_column, "") if config.job_column else record.get("报考岗位", "")

    exam_numbers = [item.get(config.exam_no_column, "") for item in sorted_records if item.get(config.exam_no_column, "")]
    subjects = _unique_join(item.get(config.subject_column, "") for item in sorted_records)
    units = _unique_join(
        _join_parts(item.get(config.unit_column, ""), item.get(config.job_column, ""), separator="")
        for item in sorted_records
    )
    first = sorted_records[0]
    context = {key: _escape(value) for key, value in first.items()}
    context.update(
        {
            "考场": _escape(room_value),
            "考点": _escape(first.get(config.site_column, "")),
            "考试科目": _escape(subjects),
            "起始考号": _escape(min(exam_numbers, default="")),
            "结束考号": _escape(max(exam_numbers, default="")),
            "单位岗位汇总": _escape(units),
            "考生人数": _escape(len(sorted_records)),
        }
    )
    context["座号"] = _escape(first.get(config.seat_column, ""))
    context["考号"] = _escape(first.get(config.exam_no_column, ""))
    context["报考单位"] = _escape(first.get(config.unit_column, ""))
    context["报考岗位"] = _escape(first.get(config.job_column, ""))
    return context, sorted_records


def render_document(
    template: PrintTemplate,
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    columns: int,
    items_per_page: int,
    *,
    sort_column: str = "",
    image_mode: str = "data",
) -> str:
    context, room_records = build_room_context(records, config, room_value, sort_column=sort_column, image_mode=image_mode)
    if template.doc_type == DOC_TYPE_SEAT:
        if str(template.settings.get("layout", "")) == "candidate_grid":
            candidate_grid = _build_candidate_grid(room_records, template.item_html, context, columns)
            context["考生列表"] = candidate_grid
            context["座次表表格"] = candidate_grid
        elif str(template.settings.get("layout", "")) == "signin":
            signin_table = _build_signin_table(room_records, room_context=context)
            context["考生列表"] = signin_table
            context["座次表表格"] = signin_table
        else:
            context["座次表表格"] = _build_seat_table(room_records, template.item_html, context, columns)
    elif template.doc_type == DOC_TYPE_DESK:
        context["桌贴列表"] = _build_desk_pages(room_records, template.item_html, context, columns, items_per_page)
    return render_html_template(template.main_html, context)


def render_document_pages(
    template: PrintTemplate,
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    columns: int,
    items_per_page: int,
    *,
    sort_column: str = "",
    page_group_column: str = "",
    image_mode: str = "data",
) -> str:
    base_config = replace(config, room_column="") if not room_value else config
    page_records = prepare_print_records(records, base_config, room_value, sort_column=sort_column)
    group_config = replace(base_config, room_column="")
    rendered_pages = [
        render_document(
            template,
            group_records,
            group_config,
            "",
            columns,
            items_per_page,
            sort_column=sort_column,
            image_mode=image_mode,
        )
        for group_records in _page_groups(page_records, page_group_column)
    ]
    return _combine_rendered_pages(rendered_pages)


def render_html_template(template_html: str, context: dict[str, str]) -> str:
    def replacer(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        return str(context.get(key, ""))

    return PLACEHOLDER_PATTERN.sub(replacer, template_html)


def _html_body_fragment(template_html: str) -> str:
    html_text = re.sub(r"(?is)<style\b[^>]*>.*?</style>", "", template_html)
    html_text = re.sub(r"(?is)<script\b[^>]*>.*?</script>", "", html_text)
    body_match = re.search(r"(?is)<body\b[^>]*>(.*?)</body>", html_text)
    if body_match:
        html_text = body_match.group(1)
    html_text = re.sub(r"(?is)<head\b[^>]*>.*?</head>", "", html_text)
    return html_text.strip()


def _html_head_fragment(template_html: str) -> str:
    head_match = re.search(r"(?is)<head\b[^>]*>(.*?)</head>", template_html)
    return head_match.group(1).strip() if head_match else ""


def _combine_rendered_pages(rendered_pages: list[str]) -> str:
    if not rendered_pages:
        return ""
    if len(rendered_pages) == 1:
        return rendered_pages[0]
    head_html = _html_head_fragment(rendered_pages[0])
    body_parts = [_html_body_fragment(page) for page in rendered_pages]
    pages_html = []
    for index, body_html in enumerate(body_parts):
        page_break = " page-break-after: always;" if index < len(body_parts) - 1 else ""
        pages_html.append(f'<div class="print-page" style="{page_break}">{body_html}</div>')
    return f"<html><head>{head_html}</head><body>{''.join(pages_html)}</body></html>"


def export_rendered_html(output_path: Path, rendered_html: str) -> None:
    output_path.write_text(rendered_html, encoding="utf-8")


def export_interview_signin_pdf(
    output_path: Path,
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    *,
    title: str,
    top_html: str = "",
    bottom_html: str = "",
    item_html: str = "",
    people_per_line: int = 5,
    sort_column: str = "",
    file_group_column: str = "",
    page_group_column: str = "",
    start_corner: str = "top_left",
    fill_direction: str = "row",
    snake: bool = False,
    photo_width: float = 52,
    font_size: float = 6.9,
    line_spacing: float = 7.45,
    section_background: str = "",
) -> list[Path]:
    from PIL import Image, ImageFile, ImageOps
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass
    font_name = "STSong-Light"
    room_records = _filtered_sorted_records(records, config, room_value, sort_column=sort_column)
    if file_group_column:
        output_paths: list[Path] = []
        for file_group_value, file_group_records in _named_record_groups(room_records, file_group_column):
            grouped_output_path = output_path.with_name(
                f"{output_path.stem}_{_safe_filename(file_group_value or '未分组')}{output_path.suffix}"
            )
            output_paths.extend(
                export_interview_signin_pdf(
                    grouped_output_path,
                    file_group_records,
                    config,
                    "",
                    title=title,
                    top_html=top_html,
                    bottom_html=bottom_html,
                    item_html=item_html,
                    people_per_line=people_per_line,
                    sort_column=sort_column,
                    file_group_column="",
                    page_group_column=page_group_column,
                    start_corner=start_corner,
                    fill_direction=fill_direction,
                    snake=snake,
                    photo_width=photo_width,
                    font_size=font_size,
                    line_spacing=line_spacing,
                    section_background=section_background,
                )
            )
        return output_paths
    photo_map = _photo_file_map(config.photo_dir)

    page_width, page_height = landscape(A4)
    margin_x = 14
    margin_top = 14
    section_default_font_size = 12.5
    top_style = _pdf_fragment_style(top_html or title, default_font_size=section_default_font_size)
    bottom_style = _pdf_fragment_style(bottom_html, default_font_size=section_default_font_size)
    if section_background.strip():
        normalized_background = _normalize_pdf_color(section_background.strip(), "#ffffff")
        top_style["background"] = normalized_background
        bottom_style["background"] = normalized_background
    top_font_size = float(top_style.get("font_size", section_default_font_size))
    bottom_font_size = float(bottom_style.get("font_size", section_default_font_size))
    top_leading = top_font_size + 4
    bottom_leading = bottom_font_size + 4
    top_lines = _pdf_fragment_lines(top_html or title, {}, font_name, top_font_size, page_width - margin_x * 2)
    bottom_lines = _pdf_fragment_lines(bottom_html, {}, font_name, bottom_font_size, page_width - margin_x * 2)
    top_height = 0 if not top_lines else max(20, len(top_lines) * top_leading + 6)
    top_gap = 7 if top_lines else 0
    bottom_height = 0 if not bottom_lines else max(16, len(bottom_lines) * bottom_leading + 6)
    bottom_gap = 7 if bottom_lines else 0
    safe_people_per_line = max(1, people_per_line)
    if fill_direction == "column":
        safe_columns = 5
        rows_per_page = safe_people_per_line
    else:
        safe_columns = safe_people_per_line
        rows_per_page = 6
    card_width = (page_width - margin_x * 2) / safe_columns
    grid_top = page_height - margin_top - top_height - top_gap
    grid_bottom = 14 + bottom_height + bottom_gap
    card_height = (grid_top - grid_bottom) / rows_per_page
    padding = 6
    safe_photo_width = 0 if photo_width <= 0 else min(max(20, photo_width), card_width * 0.55)
    photo_height = max(1, card_height - padding * 2)
    safe_font_size = max(4, font_size)
    leading = max(safe_font_size + 0.4, line_spacing)

    pdf = canvas.Canvas(str(output_path), pagesize=landscape(A4))
    pdf.setTitle(title)
    for group_records in _page_groups(room_records, page_group_column):
        for page_start in range(0, len(group_records), safe_columns * rows_per_page):
            page_records = group_records[page_start : page_start + safe_columns * rows_per_page]
            positions = _layout_positions(
                len(page_records),
                safe_columns,
                rows_per_page,
                start_corner=start_corner,
                fill_direction=fill_direction,
                snake=snake,
            )
            page_context = _pdf_group_context(page_records, config)
            page_top_lines = _pdf_fragment_lines(top_html or title, page_context, font_name, top_font_size, page_width - margin_x * 2)
            page_bottom_lines = _pdf_fragment_lines(bottom_html, page_context, font_name, bottom_font_size, page_width - margin_x * 2)
            if page_top_lines:
                title_y = page_height - margin_top - top_height
                pdf.setFillColor(colors.HexColor(str(top_style.get("background", "#f5dada"))))
                pdf.rect(margin_x + 4, title_y, page_width - (margin_x + 4) * 2, top_height, stroke=0, fill=1)
                pdf.setFillColor(colors.HexColor(str(top_style.get("color", "#c62828"))))
                pdf.setFont(font_name, top_font_size)
                line_y = title_y + top_height - top_leading
                for line in page_top_lines:
                    pdf.drawCentredString(page_width / 2, line_y, line)
                    line_y -= top_leading

            for index, record in enumerate(page_records):
                row, column = positions[index]
                x = margin_x + column * card_width
                y = grid_top - (row + 1) * card_height
                pdf.setStrokeColor(colors.HexColor("#333333"))
                pdf.setLineWidth(0.65)
                pdf.rect(x, y, card_width, card_height, stroke=1, fill=0)

                image_x = x + padding
                image_y = y + card_height - photo_height - padding
                if safe_photo_width > 0:
                    photo_path = _find_photo_file(record, config, photo_map)
                    if photo_path is not None:
                        _draw_photo(pdf, photo_path, image_x, image_y, safe_photo_width, photo_height)
                    else:
                        pdf.setStrokeColor(colors.HexColor("#aaaaaa"))
                        pdf.rect(image_x, image_y, safe_photo_width, photo_height, stroke=1, fill=0)

                text_x = x + padding if safe_photo_width <= 0 else image_x + safe_photo_width + 4
                text_y = y + card_height - padding - 3
                text_width = card_width - (text_x - x) - padding
                bottom_limit = y + 4.5
                pdf.setFillColor(colors.black)
                pdf.setFont(font_name, safe_font_size)
                lines = _candidate_pdf_lines(record, config, item_html, text_width, font_name, safe_font_size)
                for line in lines:
                    if text_y - leading < bottom_limit:
                        break
                    if "考生签字" in line:
                        pdf.drawString(text_x, text_y, line)
                        pdf.setStrokeColor(colors.HexColor("#333333"))
                        pdf.setLineWidth(0.45)
                        line_start = text_x + pdfmetrics.stringWidth(line, font_name, safe_font_size) + 4
                        if line_start < x + card_width - padding:
                            pdf.line(line_start, text_y - 1, x + card_width - padding, text_y - 1)
                    else:
                        pdf.drawString(text_x, text_y, line)
                    text_y -= leading

            if page_bottom_lines:
                bottom_y = 14
                pdf.setFillColor(colors.HexColor(str(bottom_style.get("background", "#ffffff"))))
                pdf.rect(margin_x + 4, bottom_y, page_width - (margin_x + 4) * 2, bottom_height, stroke=0, fill=1)
                pdf.setFillColor(colors.HexColor(str(bottom_style.get("color", "#111827"))))
                pdf.setFont(font_name, max(4, bottom_font_size))
                line_y = bottom_y + bottom_height - bottom_leading
                for line in page_bottom_lines:
                    pdf.drawCentredString(page_width / 2, line_y, line)
                    line_y -= bottom_leading

            pdf.showPage()
    pdf.save()
    return [output_path]


def _join_parts(left: str, right: str, separator: str = " ") -> str:
    parts = [part for part in (left.strip(), right.strip()) if part]
    return separator.join(parts)


def _unique_join(values: Iterable[str]) -> str:
    seen: list[str] = []
    for value in values:
        text = _coerce_text(value)
        if text and text not in seen:
            seen.append(text)
    return "、".join(seen)


def _filtered_sorted_records(
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    *,
    sort_column: str = "",
) -> list[dict[str, str]]:
    if config.room_column and room_value:
        room_records = [record for record in records if _coerce_text(record.get(config.room_column, "")) == room_value]
    else:
        room_records = list(records)
    if not room_records:
        raise ValueError("没有匹配到任何考生。")
    if sort_column:
        return sorted(room_records, key=lambda item: _natural_sort_key(item.get(sort_column, "")))
    return sorted(room_records, key=lambda item: _sort_key(item, config))


def prepare_print_records(
    records: list[dict[str, str]],
    config: PrintDataConfig,
    room_value: str,
    *,
    sort_column: str = "",
) -> list[dict[str, str]]:
    return _filtered_sorted_records(records, config, room_value, sort_column=sort_column)


def named_record_groups(records: list[dict[str, str]], group_column: str) -> list[tuple[str, list[dict[str, str]]]]:
    return _named_record_groups(records, group_column)


def _natural_sort_key(value: object) -> tuple[object, ...]:
    text = _coerce_text(value)
    parts = re.split(r"(\d+)", text)
    key: list[object] = []
    for part in parts:
        if part.isdigit():
            key.append((0, int(part)))
        else:
            key.append((1, part))
    return tuple(key)


def _page_groups(records: list[dict[str, str]], page_group_column: str) -> list[list[dict[str, str]]]:
    if not page_group_column:
        return [records]
    return [group_records for _, group_records in _named_record_groups(records, page_group_column)]


def _named_record_groups(records: list[dict[str, str]], group_column: str) -> list[tuple[str, list[dict[str, str]]]]:
    groups: list[list[dict[str, str]]] = []
    named_groups: list[tuple[str, list[dict[str, str]]]] = []
    group_index: dict[str, tuple[str, list[dict[str, str]]]] = {}
    for record in records:
        key = _coerce_text(record.get(group_column, ""))
        if key not in group_index:
            item = (key, [])
            group_index[key] = item
            named_groups.append(item)
        group_index[key][1].append(record)
    return named_groups


def _safe_filename(value: str) -> str:
    safe_value = re.sub(r'[\\/:*?"<>|]+', "_", _coerce_text(value))
    safe_value = safe_value.strip().strip(".")
    return safe_value or "未分组"


def _layout_positions(
    count: int,
    columns: int,
    rows: int,
    *,
    start_corner: str,
    fill_direction: str,
    snake: bool,
) -> list[tuple[int, int]]:
    top_first = not start_corner.startswith("bottom")
    left_first = not start_corner.endswith("right")
    row_order = list(range(rows)) if top_first else list(reversed(range(rows)))
    column_order = list(range(columns)) if left_first else list(reversed(range(columns)))
    positions: list[tuple[int, int]] = []
    if fill_direction == "column":
        for column_index, column in enumerate(column_order):
            current_rows = row_order
            if snake and column_index % 2 == 1:
                current_rows = list(reversed(row_order))
            for row in current_rows:
                positions.append((row, column))
                if len(positions) >= count:
                    return positions
    else:
        for row_index, row in enumerate(row_order):
            current_columns = column_order
            if snake and row_index % 2 == 1:
                current_columns = list(reversed(column_order))
            for column in current_columns:
                positions.append((row, column))
                if len(positions) >= count:
                    return positions
    return positions


def _photo_file_map(photo_dir: Path | None) -> dict[str, Path]:
    if photo_dir is None or not photo_dir.exists():
        return {}
    photo_map: dict[str, Path] = {}
    for file_path in sorted(photo_dir.rglob("*")):
        if file_path.name.startswith("._"):
            continue
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_IMAGE_SUFFIXES:
            photo_map[_normalize_match_text(file_path.stem)] = file_path
    return photo_map


def _find_photo_file(record: dict[str, str], config: PrintDataConfig, photo_map: dict[str, Path]) -> Path | None:
    if not config.photo_match_column:
        return None
    match_value = _normalize_match_text(record.get(config.photo_match_column, ""))
    if not match_value:
        return None
    if match_value in photo_map:
        return photo_map[match_value]
    for stem, file_path in photo_map.items():
        if stem.startswith(match_value) or match_value in stem:
            return file_path
    return None


def _draw_photo(pdf, photo_path: Path, x: float, y: float, width: float, height: float) -> None:
    from PIL import Image, ImageOps
    from reportlab.lib.utils import ImageReader

    with Image.open(photo_path) as image:
        image = ImageOps.exif_transpose(image)
        image = ImageOps.fit(image.convert("RGB"), (int(width * 3), int(height * 3)), method=Image.Resampling.LANCZOS)
        pdf.drawImage(ImageReader(image), x, y, width=width, height=height, preserveAspectRatio=False, mask=None)


def _wrap_pdf_text(text: str, max_width: float, font_name: str, font_size: float) -> list[str]:
    from reportlab.pdfbase import pdfmetrics

    lines: list[str] = []
    current = ""
    for char in _coerce_text(text):
        candidate = current + char
        if pdfmetrics.stringWidth(candidate, font_name, font_size) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines or [""]


def _candidate_pdf_lines(
    record: dict[str, str],
    config: PrintDataConfig,
    item_html: str,
    max_width: float,
    font_name: str,
    font_size: float,
) -> list[str]:
    if item_html.strip():
        context = _merge_context(record, {})
        context["报考单位"] = _escape(_record_unit(record, config))
        context["报考岗位"] = _escape(_record_job(record, config))
        rendered = render_html_template(_html_body_fragment(item_html), context)
        rendered = re.sub(r"<img\b[^>]*>", "", rendered, flags=re.IGNORECASE)
        rendered = re.sub(r"(?i)<br\s*/?>", "\n", rendered)
        rendered = re.sub(r"(?i)</(div|p|tr|li|td)>", "\n", rendered)
        rendered = re.sub(r"<[^>]+>", "", rendered)
        raw_lines = [html.unescape(line).strip() for line in rendered.splitlines()]
        raw_lines = [line for line in raw_lines if line]
    else:
        raw_lines = [
            f"姓名:{_coerce_text(record.get('姓名', ''))}",
            "身份证号:",
            _coerce_text(record.get("身份证号", "")),
            f"报考单位:{_record_unit(record, config)}",
            f"报考职位:{_record_job(record, config)}",
        ]
    lines: list[str] = []
    for raw_line in raw_lines:
        if "考生签字" in raw_line:
            lines.append(raw_line)
        else:
            lines.extend(_wrap_pdf_text(raw_line, max_width, font_name, font_size))
    return lines


def _pdf_group_context(records: list[dict[str, str]], config: PrintDataConfig) -> dict[str, str]:
    if not records:
        return {}
    first = records[0]
    exam_numbers = [item.get(config.exam_no_column, "") for item in records if config.exam_no_column and item.get(config.exam_no_column, "")]
    subjects = _unique_join(item.get(config.subject_column, "") for item in records if config.subject_column)
    units = _unique_join(
        _join_parts(item.get(config.unit_column, ""), item.get(config.job_column, ""), separator="")
        for item in records
    )
    context = {key: _escape(value) for key, value in first.items()}
    context.update(
        {
            "考场": _escape(first.get(config.room_column, "")),
            "考点": _escape(first.get(config.site_column, "")),
            "考试科目": _escape(subjects),
            "起始考号": _escape(min(exam_numbers, default="")),
            "结束考号": _escape(max(exam_numbers, default="")),
            "单位岗位汇总": _escape(units),
            "考生人数": _escape(len(records)),
            "座号": _escape(first.get(config.seat_column, "")),
            "考号": _escape(first.get(config.exam_no_column, "")),
            "报考单位": _escape(_record_unit(first, config)),
            "报考岗位": _escape(_record_job(first, config)),
        }
    )
    return context


def _pdf_fragment_lines(
    fragment_html: str,
    context: dict[str, str],
    font_name: str,
    font_size: float,
    max_width: float,
) -> list[str]:
    fragment = _html_body_fragment(fragment_html)
    rendered = render_html_template(fragment, context)
    rendered = re.sub(r"(?i)<br\s*/?>", "\n", rendered)
    rendered = re.sub(r"(?i)</(div|p|tr|li|td|h[1-6])>", "\n", rendered)
    rendered = re.sub(r"<[^>]+>", "", rendered)
    raw_lines = [html.unescape(line).strip() for line in rendered.splitlines()]
    lines: list[str] = []
    for raw_line in raw_lines:
        if not raw_line:
            continue
        lines.extend(_wrap_pdf_text(raw_line, max_width, font_name, font_size))
    return lines


def _pdf_fragment_style(fragment_html: str, *, default_font_size: float) -> dict[str, object]:
    style_match = re.search(r'(?is)style\s*=\s*["\']([^"\']+)["\']', fragment_html)
    css = style_match.group(1) if style_match else ""
    style: dict[str, object] = {"font_size": default_font_size}
    for part in css.split(";"):
        if ":" not in part:
            continue
        name, value = part.split(":", 1)
        name = name.strip().lower()
        value = value.strip()
        if name == "color":
            style["color"] = _normalize_pdf_color(value, "#111827")
        elif name in {"background", "background-color"}:
            style["background"] = _normalize_pdf_color(value, "#f5dada")
        elif name == "font-size":
            style["font_size"] = _css_point_size(value, default_font_size)
    if "color" not in style:
        style["color"] = "#111827"
    return style


def _css_point_size(value: str, default: float) -> float:
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)", value)
    if not match:
        return default
    number = float(match.group(1))
    return number


def _normalize_pdf_color(value: str, default: str) -> str:
    text = value.strip()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", text):
        return text
    if re.fullmatch(r"#[0-9a-fA-F]{3}", text):
        return "#" + "".join(char * 2 for char in text[1:])
    named = {
        "black": "#111827",
        "red": "#c62828",
        "white": "#ffffff",
        "gray": "#6b7280",
        "grey": "#6b7280",
        "blue": "#1d4ed8",
    }
    return named.get(text.lower(), default)


def _record_unit(record: dict[str, str], config: PrintDataConfig) -> str:
    return _coerce_text(record.get("报考单位") or record.get(config.unit_column, "") or record.get("单位", ""))


def _record_job(record: dict[str, str], config: PrintDataConfig) -> str:
    return _coerce_text(record.get("报考岗位") or record.get(config.job_column, "") or record.get("岗位", ""))


def _match_photo(record: dict[str, str], config: PrintDataConfig, *, image_mode: str = "data") -> str:
    if config.photo_dir is None or not config.photo_dir.exists() or not config.photo_match_column:
        return ""
    match_value = _normalize_match_text(record.get(config.photo_match_column, ""))
    if not match_value:
        return ""
    for file_path in sorted(config.photo_dir.rglob("*")):
        if not file_path.is_file() or file_path.suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES:
            continue
        stem = _normalize_match_text(file_path.stem)
        if stem == match_value or stem.startswith(match_value) or match_value in stem:
            return _to_file_url(file_path) if image_mode == "file" else _to_image_data_uri(file_path)
    return ""


def _normalize_match_text(value: object) -> str:
    text = _coerce_text(value).upper()
    return re.sub(r"[^0-9A-Z]", "", text)


def _sort_key(record: dict[str, str], config: PrintDataConfig) -> tuple[int, str, str]:
    seat_text = _coerce_text(record.get(config.seat_column, ""))
    seat_number = int(seat_text) if seat_text.isdigit() else math.inf
    return seat_number, _coerce_text(record.get(config.exam_no_column, "")), _coerce_text(record.get("姓名", ""))


def _merge_context(record: dict[str, str], room_context: dict[str, str]) -> dict[str, str]:
    merged = dict(room_context)
    merged.update({key: _escape(value) for key, value in record.items()})
    merged["座号"] = _escape(record.get("座号") or record.get("seat", ""))
    merged["考号"] = _escape(record.get("考号") or record.get("准考证号") or record.get("exam_no", ""))
    merged["报考单位"] = _escape(record.get("报考单位") or record.get("单位") or "")
    merged["报考岗位"] = _escape(record.get("报考岗位") or record.get("岗位") or "")
    if record.get("照片"):
        merged["照片"] = record["照片"]
        merged["照片标签"] = f'<img class="photo" src="{record["照片"]}" />'
    else:
        merged["照片标签"] = ""
    return merged


def _build_seat_table(records: list[dict[str, str]], item_html: str, room_context: dict[str, str], columns: int) -> str:
    safe_columns = max(1, columns)
    item_fragment = _html_body_fragment(item_html)
    rows: list[str] = []
    cells: list[str] = []
    for record in records:
        merged_context = _merge_context(record, room_context)
        cells.append(f"<td>{render_html_template(item_fragment, merged_context)}</td>")
        if len(cells) == safe_columns:
            rows.append("<tr>" + "".join(cells) + "</tr>")
            cells = []
    if cells:
        while len(cells) < safe_columns:
            cells.append("<td></td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table class="seat-table">' + "".join(rows) + "</table>"


def _build_candidate_grid(records: list[dict[str, str]], item_html: str, room_context: dict[str, str], columns: int) -> str:
    safe_columns = max(1, columns)
    item_fragment = _html_body_fragment(item_html)
    rows: list[str] = []
    cells: list[str] = []
    for record in records:
        merged_context = _merge_context(record, room_context)
        cells.append(
            '<td style="border:1px solid #444;padding:0;height:94px;vertical-align:top;">'
            f"{render_html_template(item_fragment, merged_context)}"
            "</td>"
        )
        if len(cells) == safe_columns:
            rows.append("<tr>" + "".join(cells) + "</tr>")
            cells = []
    if cells:
        while len(cells) < safe_columns:
            cells.append('<td style="border:1px solid #444;padding:0;height:94px;"></td>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table class="candidate-grid">' + "".join(rows) + "</table>"


def _build_signin_table(records: list[dict[str, str]], room_context: dict[str, str]) -> str:
    header = """
<tr>
  <th style="width: 7%;">序号</th>
  <th style="width: 9%;">照片</th>
  <th style="width: 9%;">姓名</th>
  <th style="width: 11%;">考号</th>
  <th style="width: 15%;">报考单位</th>
  <th style="width: 15%;">报考岗位</th>
  <th style="width: 18%;">身份证号</th>
  <th style="width: 9%;">报到时间</th>
  <th style="width: 7%;">签到</th>
</tr>
""".strip()
    rows: list[str] = []
    for index, record in enumerate(records, start=1):
        merged = _merge_context(record, room_context)
        photo_url = record.get("照片", "")
        photo_cell = f'<img class="signin-photo" src="{photo_url}" />' if photo_url else ""
        rows.append(
            "<tr>"
            f"<td>{index}</td>"
            f"<td>{photo_cell}</td>"
            f"<td>{merged.get('姓名', '')}</td>"
            f"<td>{merged.get('考号', '')}</td>"
            f"<td>{merged.get('报考单位', '')}</td>"
            f"<td>{merged.get('报考岗位', '')}</td>"
            f"<td>{merged.get('身份证号', '')}</td>"
            f"<td>{merged.get('面试报到时间', '')}</td>"
            '<td class="signin-signature"></td>'
            "</tr>"
        )
    return '<table class="signin-table">' + header + "".join(rows) + "</table>"


def _build_desk_pages(
    records: list[dict[str, str]],
    item_html: str,
    room_context: dict[str, str],
    columns: int,
    items_per_page: int,
) -> str:
    safe_columns = max(1, columns)
    safe_items = max(1, items_per_page)
    item_fragment = _html_body_fragment(item_html)
    pages: list[str] = []
    for page_start in range(0, len(records), safe_items):
        chunk = records[page_start : page_start + safe_items]
        rows: list[str] = []
        cells: list[str] = []
        for record in chunk:
            merged_context = _merge_context(record, room_context)
            cells.append(f"<td>{render_html_template(item_fragment, merged_context)}</td>")
            if len(cells) == safe_columns:
                rows.append("<tr>" + "".join(cells) + "</tr>")
                cells = []
        if cells:
            while len(cells) < safe_columns:
                cells.append("<td></td>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        pages.append(f'<div class="page"><table class="desk-grid">{"".join(rows)}</table></div>')
    return "".join(pages)
