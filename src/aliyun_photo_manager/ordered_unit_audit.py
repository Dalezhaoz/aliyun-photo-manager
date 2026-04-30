from __future__ import annotations

from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")
ISSUE_FILL = PatternFill(fill_type="solid", fgColor="FDE2E1")
ISSUE_FONT = Font(color="C00000", bold=True)


@dataclass(frozen=True)
class OrderedNameIssue:
    issue_type: str
    first_row: int
    second_row: int
    first_name: str
    second_name: str
    description: str


@dataclass(frozen=True)
class OrderedNameAuditResult:
    source_path: Path
    target_column: str
    total_rows: int
    issues: list[OrderedNameIssue]


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_excel_headers(excel_path: Path) -> list[str]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [_text(value) for value in header_row if _text(value)]
    finally:
        workbook.close()


def _normalize_name(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    replacements = (
        ("（", "("),
        ("）", ")"),
        (" ", ""),
        ("\u3000", ""),
        ("委员会", "委"),
        ("卫生和健康委员会", "卫健委"),
        ("卫生健康委员会", "卫健委"),
        ("卫生和健康委", "卫健委"),
        ("卫生健康委", "卫健委"),
        ("教育和体育局", "教体局"),
        ("教育体育局", "教体局"),
        ("人力资源和社会保障局", "人社局"),
        ("人力资源社会保障局", "人社局"),
    )
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def _is_similar(left: str, right: str) -> bool:
    if not left or not right:
        return False
    left_norm = _normalize_name(left)
    right_norm = _normalize_name(right)
    if not left_norm or not right_norm:
        return False
    if left_norm == right_norm:
        return True
    if left_norm in right_norm or right_norm in left_norm:
        return True
    return SequenceMatcher(None, left_norm, right_norm).ratio() >= 0.82


def run_ordered_name_audit(excel_path: Path, target_column: str) -> OrderedNameAuditResult:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, ())
        headers = [_text(value) for value in header_row]
        try:
            column_index = headers.index(target_column)
        except ValueError as exc:
            raise ValueError(f"未找到列：{target_column}") from exc

        ordered_items: list[tuple[int, str]] = []
        for row_number, row in enumerate(rows, start=2):
            value = _text(row[column_index] if column_index < len(row) else "")
            if value:
                ordered_items.append((row_number, value))
    finally:
        workbook.close()

    issues: list[OrderedNameIssue] = []

    index = 1
    while index < len(ordered_items):
        prev_row, prev_name = ordered_items[index - 1]
        curr_row, curr_name = ordered_items[index]
        if not _is_similar(prev_name, curr_name):
            index += 1
            continue

        group_start_row = prev_row
        group_start_name = prev_name
        group_end_row = curr_row
        group_end_name = curr_name

        while index + 1 < len(ordered_items):
            next_row, next_name = ordered_items[index + 1]
            if not _is_similar(group_end_name, next_name):
                break
            group_end_row = next_row
            group_end_name = next_name
            index += 1

        issues.append(
            OrderedNameIssue(
                issue_type="上下相似",
                first_row=group_start_row,
                second_row=group_end_row,
                first_name=group_start_name,
                second_name=group_end_name,
                description=(
                    f"第 {group_start_row} 行到第 {group_end_row} 行连续出现相似名称，"
                    "建议人工核对是否应统一。"
                ),
            )
        )
        index += 1

    seen_segments: list[tuple[int, str, int]] = []
    current_name = ""
    current_start = 0
    current_end = 0

    for row_number, name in ordered_items:
        if not current_name:
            current_name = name
            current_start = row_number
            current_end = row_number
            continue

        if _is_similar(current_name, name):
            current_end = row_number
            continue

        seen_segments.append((current_start, current_name, current_end))
        for previous_start, previous_name, previous_end in seen_segments[:-1]:
            if _is_similar(previous_name, name):
                issues.append(
                    OrderedNameIssue(
                        issue_type="分段重复",
                        first_row=previous_start,
                        second_row=row_number,
                        first_name=previous_name,
                        second_name=name,
                        description=(
                            f"“{name}”前面已在第 {previous_start}-{previous_end} 行出现过，"
                            f"中间夹了别的内容后又在第 {row_number} 行重新出现，建议核对顺序是否应连续。"
                        ),
                    )
                )
                break

        current_name = name
        current_start = row_number
        current_end = row_number

    return OrderedNameAuditResult(
        source_path=excel_path,
        target_column=target_column,
        total_rows=len(ordered_items),
        issues=issues,
    )


def export_ordered_name_audit(result: OrderedNameAuditResult, output_path: Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "核对摘要"
    summary.append(["源文件", str(result.source_path)])
    summary.append(["目标列", result.target_column])
    summary.append(["参与核对行数", result.total_rows])
    summary.append(["问题条数", len(result.issues)])
    summary.append([])
    summary.append(["说明", "严格按原表顺序核对当前列，不会先排序。"])

    detail = workbook.create_sheet("需核对清单")
    headers = ["问题类型", "上一处行号", "当前行号", "上一处名称", "当前名称", "说明"]
    detail.append(headers)
    for cell in detail[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    for issue in result.issues:
        detail.append(
            [
                issue.issue_type,
                issue.first_row,
                issue.second_row,
                issue.first_name,
                issue.second_name,
                issue.description,
            ]
        )
        issue_cell = detail.cell(detail.max_row, 1)
        issue_cell.fill = ISSUE_FILL
        issue_cell.font = ISSUE_FONT

    widths = {"A": 16, "B": 14, "C": 14, "D": 28, "E": 28, "F": 60}
    for column, width in widths.items():
        detail.column_dimensions[column].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
