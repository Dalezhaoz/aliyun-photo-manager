from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
import json
from pathlib import Path
from typing import Callable, Iterable, List, Optional

from openpyxl import Workbook


Logger = Optional[Callable[[str], None]]


@dataclass
class StageServerConfig:
    name: str
    host: str
    port: int
    username: str
    password: str
    enabled: bool = True


@dataclass
class ProjectStageRecord:
    server_name: str
    database_name: str
    project_name: str
    stage_name: str
    start_time: datetime
    end_time: datetime
    status: str


@dataclass
class ProjectStageSummary:
    records: List[ProjectStageRecord]
    enabled_servers: int
    visited_databases: int
    matched_databases: int
    ongoing_count: int
    upcoming_count: int


def summary_to_dict(summary: ProjectStageSummary) -> dict:
    return {
        "records": [
            {
                **asdict(item),
                "start_time": item.start_time.isoformat(),
                "end_time": item.end_time.isoformat(),
            }
            for item in summary.records
        ],
        "enabled_servers": summary.enabled_servers,
        "visited_databases": summary.visited_databases,
        "matched_databases": summary.matched_databases,
        "ongoing_count": summary.ongoing_count,
        "upcoming_count": summary.upcoming_count,
    }


def summary_from_dict(data: dict) -> ProjectStageSummary:
    return ProjectStageSummary(
        records=[
            ProjectStageRecord(
                server_name=str(item.get("server_name", "")),
                database_name=str(item.get("database_name", "")),
                project_name=str(item.get("project_name", "")),
                stage_name=str(item.get("stage_name", "")),
                start_time=datetime.fromisoformat(str(item.get("start_time", ""))),
                end_time=datetime.fromisoformat(str(item.get("end_time", ""))),
                status=str(item.get("status", "")),
            )
            for item in data.get("records", [])
        ],
        enabled_servers=int(data.get("enabled_servers", 0)),
        visited_databases=int(data.get("visited_databases", 0)),
        matched_databases=int(data.get("matched_databases", 0)),
        ongoing_count=int(data.get("ongoing_count", 0)),
        upcoming_count=int(data.get("upcoming_count", 0)),
    )


def _log(logger: Logger, message: str) -> None:
    if logger is not None:
        logger(message)


def _load_pyodbc():
    try:
        import pyodbc  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "未安装 pyodbc，无法连接 SQL Server。请先执行 `pip install -r requirements.txt`。"
        ) from exc
    return pyodbc


def _pick_driver(pyodbc) -> str:
    preferred = [
        "ODBC Driver 18 for SQL Server",
        "ODBC Driver 17 for SQL Server",
        "SQL Server",
    ]
    installed = list(pyodbc.drivers())
    for name in preferred:
        if name in installed:
            return name
    if installed:
        return installed[-1]
    raise RuntimeError("未找到 SQL Server ODBC 驱动，请先安装对应驱动。")


def _ordered_drivers(pyodbc) -> list[str]:
    preferred = [
        "ODBC Driver 17 for SQL Server",
        "ODBC Driver 18 for SQL Server",
        "SQL Server",
    ]
    installed = list(pyodbc.drivers())
    ordered = [name for name in preferred if name in installed]
    for name in installed:
        if name not in ordered:
            ordered.append(name)
    if not ordered:
        raise RuntimeError("未找到 SQL Server ODBC 驱动，请先安装对应驱动。")
    return ordered


def _status_from_times(now: datetime, start_time: datetime, end_time: datetime) -> str:
    if now < start_time:
        return "即将开始"
    if now > end_time:
        return "已经结束"
    return "正在进行"


def _status_allowed(status_filter: str, status: str) -> bool:
    if status_filter == "全部":
        return True
    if status_filter == "只看正在进行":
        return status == "正在进行"
    if status_filter == "只看即将开始":
        return status == "即将开始"
    return status in {"正在进行", "即将开始"}


def _escape_db_name(name: str) -> str:
    return name.replace("]", "]]")


def _tables_exist(cursor, database_name: str) -> bool:
    db_name = _escape_db_name(database_name)
    sql = f"""
    SELECT COUNT(*) 
    FROM [{db_name}].sys.tables
    WHERE name IN ('EI_ExamTreeDesc', 'web_SR_CodeItem', 'WEB_SR_SetTime')
    """
    cursor.execute(sql)
    count = int(cursor.fetchone()[0])
    return count == 3


def _iter_database_names(cursor) -> Iterable[str]:
    cursor.execute(
        """
        SELECT name
        FROM sys.databases
        WHERE state_desc = 'ONLINE'
          AND name NOT IN ('master', 'model', 'msdb', 'tempdb')
        ORDER BY name
        """
    )
    for row in cursor.fetchall():
        yield str(row[0])


def _fetch_records_from_database(
    cursor,
    server_name: str,
    database_name: str,
    now: datetime,
    status_filter: str,
    stage_keyword: str,
    project_keyword: str,
) -> List[ProjectStageRecord]:
    db_name = _escape_db_name(database_name)
    sql = f"""
    SELECT
        A.NAME,
        B.Description,
        C.KDate,
        C.ZDate
    FROM [{db_name}].[dbo].[EI_ExamTreeDesc] A
    JOIN [{db_name}].[dbo].[WEB_SR_SetTime] C
        ON A.Code = C.ExamSort
    JOIN [{db_name}].[dbo].[web_SR_CodeItem] B
        ON B.Codeid = 'WT'
       AND B.Code = C.Kind
    WHERE A.CodeLen = '2'
    ORDER BY C.KDate ASC
    """
    cursor.execute(sql)
    records: List[ProjectStageRecord] = []
    stage_keyword_lower = stage_keyword.strip().lower()
    project_keyword_lower = project_keyword.strip().lower()
    for row in cursor.fetchall():
        project_name = str(row[0] or "").strip()
        stage_name = str(row[1] or "").strip()
        start_time = row[2]
        end_time = row[3]
        if not project_name or not stage_name or start_time is None or end_time is None:
            continue
        if stage_keyword_lower and stage_keyword_lower not in stage_name.lower():
            continue
        if project_keyword_lower and project_keyword_lower not in project_name.lower():
            continue
        status = _status_from_times(now, start_time, end_time)
        if not _status_allowed(status_filter, status):
            continue
        records.append(
            ProjectStageRecord(
                server_name=server_name,
                database_name=database_name,
                project_name=project_name,
                stage_name=stage_name,
                start_time=start_time,
                end_time=end_time,
                status=status,
            )
        )
    return records


def query_project_stages(
    servers: List[StageServerConfig],
    status_filter: str,
    stage_keyword: str = "",
    project_keyword: str = "",
    logger: Logger = None,
) -> ProjectStageSummary:
    pyodbc = _load_pyodbc()
    enabled_servers = [server for server in servers if server.enabled]
    if not enabled_servers:
        raise ValueError("请至少启用一台数据库服务器。")

    now = datetime.now()
    all_records: List[ProjectStageRecord] = []
    visited_databases = 0
    matched_databases = 0

    for server in enabled_servers:
        _log(logger, f"开始连接服务器：{server.name} ({server.host}:{server.port})")
        connection = None
        last_error: Exception | None = None
        for driver in _ordered_drivers(pyodbc):
            for encrypt_option in ("no", "optional"):
                connection_string = (
                    f"DRIVER={{{driver}}};"
                    f"SERVER={server.host},{server.port};"
                    f"DATABASE=master;"
                    f"UID={server.username};"
                    f"PWD={server.password};"
                    f"Encrypt={encrypt_option};"
                    "TrustServerCertificate=yes;"
                    "Connection Timeout=5;"
                )
                try:
                    connection = pyodbc.connect(connection_string, timeout=5)
                    _log(logger, f"服务器 {server.name} 使用驱动 {driver} 连接成功。")
                    break
                except Exception as exc:
                    last_error = exc
                    continue
            if connection is not None:
                break

        if connection is None:
            error_text = str(last_error) if last_error is not None else "未知错误"
            if "unsupported protocol" in error_text.lower():
                raise RuntimeError(
                    f"连接服务器失败：{server.name}。当前 SQL Server 使用的 SSL/TLS 协议过旧，"
                    "与本机 ODBC 驱动不兼容。建议优先在 Windows 上连接该服务器，或升级服务器 TLS 配置。"
                ) from last_error
            raise RuntimeError(f"连接服务器失败：{server.name}，原因：{error_text}") from last_error

        with connection:
            cursor = connection.cursor()
            for database_name in _iter_database_names(cursor):
                visited_databases += 1
                try:
                    if not _tables_exist(cursor, database_name):
                        continue
                except Exception:
                    continue
                matched_databases += 1
                _log(logger, f"检查数据库：{server.name}/{database_name}")
                try:
                    all_records.extend(
                        _fetch_records_from_database(
                            cursor,
                            server_name=server.name,
                            database_name=database_name,
                            now=now,
                            status_filter=status_filter,
                            stage_keyword=stage_keyword,
                            project_keyword=project_keyword,
                        )
                    )
                except Exception as exc:
                    _log(logger, f"数据库 {server.name}/{database_name} 查询失败：{exc}")

    ongoing_count = sum(1 for item in all_records if item.status == "正在进行")
    upcoming_count = sum(1 for item in all_records if item.status == "即将开始")
    all_records.sort(
        key=lambda item: (
            0 if item.status == "正在进行" else 1 if item.status == "即将开始" else 2,
            item.start_time,
            item.server_name,
            item.database_name,
            item.project_name,
        )
    )
    return ProjectStageSummary(
        records=all_records,
        enabled_servers=len(enabled_servers),
        visited_databases=visited_databases,
        matched_databases=matched_databases,
        ongoing_count=ongoing_count,
        upcoming_count=upcoming_count,
    )


def export_project_stages(summary: ProjectStageSummary, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目阶段汇总"
    sheet.append(
        [
            "服务器",
            "数据库",
            "项目名称",
            "阶段名称",
            "开始时间",
            "结束时间",
            "当前状态",
        ]
    )
    for item in summary.records:
        sheet.append(
            [
                item.server_name,
                item.database_name,
                item.project_name,
                item.stage_name,
                item.start_time.strftime("%Y-%m-%d %H:%M:%S"),
                item.end_time.strftime("%Y-%m-%d %H:%M:%S"),
                item.status,
            ]
        )
    info_sheet = workbook.create_sheet("统计")
    info_sheet.append(["启用服务器数", summary.enabled_servers])
    info_sheet.append(["遍历数据库数", summary.visited_databases])
    info_sheet.append(["匹配数据库数", summary.matched_databases])
    info_sheet.append(["正在进行", summary.ongoing_count])
    info_sheet.append(["即将开始", summary.upcoming_count])
    workbook.save(output_path)
    return output_path


def dump_status_query_payload(
    servers: List[StageServerConfig],
    status_filter: str,
    stage_keyword: str,
    project_keyword: str,
    output_path: Path,
) -> Path:
    payload = {
        "servers": [asdict(item) for item in servers],
        "status_filter": status_filter,
        "stage_keyword": stage_keyword,
        "project_keyword": project_keyword,
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return output_path
