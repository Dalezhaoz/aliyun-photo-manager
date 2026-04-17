from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


REQUIRED_COLUMNS = [
    "地市",
    "地市代码",
    "主管部门",
    "主管部门编码",
    "报考单位",
    "报考单位编码",
    "报考岗位",
    "报考岗位编码",
]

CODE_LENGTHS = {
    "地市代码": 2,
    "主管部门编码": 3,
    "报考单位编码": 3,
    "报考岗位编码": 2,
}


@dataclass(frozen=True)
class AuditIssue:
    issue_type: str
    level: str
    excel_rows: str
    city: str
    city_code: str
    department: str
    department_code: str
    unit: str
    unit_code: str
    job: str
    job_code: str
    description: str


@dataclass(frozen=True)
class AuditResult:
    source_path: Path
    total_rows: int
    issues: list[AuditIssue]

    @property
    def issue_count(self) -> int:
        return len(self.issues)


LEVEL_ORDER = {
    "基础": 0,
    "地市": 1,
    "主管部门": 2,
    "报考单位": 3,
    "报考岗位": 4,
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")
ISSUE_FILL = PatternFill(fill_type="solid", fgColor="FDE2E1")
ISSUE_FONT = Font(color="C00000", bold=True)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_code(value: object, width: int) -> str:
    raw = _text(value)
    if not raw:
        return ""
    if raw.isdigit():
        return raw.zfill(width)
    return raw


def _load_records(excel_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, ())
        headers = [_text(value) for value in header_row]
        records: list[dict[str, str]] = []
        for row_index, row in enumerate(rows, start=2):
            record: dict[str, str] = {"_row": str(row_index)}
            has_value = False
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = row[index] if index < len(row) else ""
                if header in CODE_LENGTHS:
                    cell_value = _normalize_code(value, CODE_LENGTHS[header])
                else:
                    cell_value = _text(value)
                record[header] = cell_value
                has_value = has_value or bool(cell_value)
            if has_value:
                records.append(record)
        return records
    finally:
        workbook.close()


def load_job_audit_headers(excel_path: Path) -> list[str]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [_text(value) for value in first_row if _text(value)]
    finally:
        workbook.close()


def validate_required_columns(headers: Iterable[str]) -> list[str]:
    header_set = {header.strip() for header in headers if header and header.strip()}
    return [column for column in REQUIRED_COLUMNS if column not in header_set]


def run_job_code_audit(excel_path: Path) -> AuditResult:
    records = _load_records(excel_path)
    issues: list[AuditIssue] = []
    if not records:
        return AuditResult(source_path=excel_path, total_rows=0, issues=issues)

    def add_issue(
        issue_type: str,
        level: str,
        rows: Iterable[str],
        description: str,
        *,
        city: str = "",
        city_code: str = "",
        department: str = "",
        department_code: str = "",
        unit: str = "",
        unit_code: str = "",
        job: str = "",
        job_code: str = "",
    ) -> None:
        row_values = sorted({row for row in rows if row}, key=lambda item: int(item))
        issues.append(
            AuditIssue(
                issue_type=issue_type,
                level=level,
                excel_rows="、".join(row_values),
                city=city,
                city_code=city_code,
                department=department,
                department_code=department_code,
                unit=unit,
                unit_code=unit_code,
                job=job,
                job_code=job_code,
                description=description,
            )
        )

    def grouped_rows(data: list[dict[str, str]], *keys: str) -> dict[tuple[str, ...], list[dict[str, str]]]:
        groups: dict[tuple[str, ...], list[dict[str, str]]] = defaultdict(list)
        for item in data:
            groups[tuple(item.get(key, "") for key in keys)].append(item)
        return groups

    def detect_mapping_conflicts(
        data: list[dict[str, str]],
        *,
        scope_keys: tuple[str, ...],
        code_key: str,
        name_key: str,
        level: str,
    ) -> None:
        by_code: dict[tuple[str, ...], dict[str, list[dict[str, str]]]] = defaultdict(lambda: defaultdict(list))
        by_name: dict[tuple[str, ...], dict[str, list[dict[str, str]]]] = defaultdict(lambda: defaultdict(list))
        for item in data:
            scope = tuple(item.get(key, "") for key in scope_keys)
            code = item.get(code_key, "")
            name = item.get(name_key, "")
            if code:
                by_code[scope][code].append(item)
            if name:
                by_name[scope][name].append(item)

        for scope, code_map in by_code.items():
            for code, rows in code_map.items():
                names = {row.get(name_key, "") for row in rows}
                if len(names) > 1:
                    sample = rows[0]
                    add_issue(
                        "编码映射冲突",
                        level,
                        [row["_row"] for row in rows],
                        f"相同的{code_key} {code} 绑定了多个{name_key}：{', '.join(sorted(names))}",
                        city=sample.get("地市", ""),
                        city_code=sample.get("地市代码", ""),
                        department=sample.get("主管部门", ""),
                        department_code=sample.get("主管部门编码", ""),
                        unit=sample.get("报考单位", ""),
                        unit_code=sample.get("报考单位编码", ""),
                        job=sample.get("报考岗位", ""),
                        job_code=sample.get("报考岗位编码", ""),
                    )

        for scope, name_map in by_name.items():
            for name, rows in name_map.items():
                codes = {row.get(code_key, "") for row in rows}
                if len(codes) > 1:
                    sample = rows[0]
                    add_issue(
                        "名称映射冲突",
                        level,
                        [row["_row"] for row in rows],
                        f"相同的{name_key} {name} 对应了多个{code_key}：{', '.join(sorted(codes))}",
                        city=sample.get("地市", ""),
                        city_code=sample.get("地市代码", ""),
                        department=sample.get("主管部门", ""),
                        department_code=sample.get("主管部门编码", ""),
                        unit=sample.get("报考单位", ""),
                        unit_code=sample.get("报考单位编码", ""),
                        job=sample.get("报考岗位", ""),
                        job_code=sample.get("报考岗位编码", ""),
                    )

    def detect_sequence_issues(
        data: list[dict[str, str]],
        *,
        parent_keys: tuple[str, ...],
        child_code_key: str,
        level: str,
        width: int,
    ) -> None:
        groups = grouped_rows(data, *parent_keys)
        for parent, rows in groups.items():
            codes = sorted({row.get(child_code_key, "") for row in rows if row.get(child_code_key, "").isdigit()})
            if not codes:
                continue
            numbers = sorted({int(code) for code in codes})
            expected = list(range(1, max(numbers) + 1))
            missing = [number for number in expected if number not in numbers]
            sample = rows[0]
            if numbers[0] != 1:
                add_issue(
                    "顺延起始错误",
                    level,
                    [row["_row"] for row in rows],
                    f"{child_code_key} 应从 {'0' * (width - 1)}1 开始，当前最小值是 {str(numbers[0]).zfill(width)}",
                    city=sample.get("地市", ""),
                    city_code=sample.get("地市代码", ""),
                    department=sample.get("主管部门", ""),
                    department_code=sample.get("主管部门编码", ""),
                    unit=sample.get("报考单位", ""),
                    unit_code=sample.get("报考单位编码", ""),
                    job=sample.get("报考岗位", ""),
                    job_code=sample.get("报考岗位编码", ""),
                )
    for record in records:
        row_no = record["_row"]
        for column in REQUIRED_COLUMNS:
            value = record.get(column, "")
            if column in {"主管部门", "报考单位", "报考岗位"}:
                continue
            if not value:
                add_issue("缺少必填值", "基础", [row_no], f"{column} 不能为空")
        for column, width in CODE_LENGTHS.items():
            value = record.get(column, "")
            if not value:
                continue
            if not value.isdigit() or len(value) != width:
                add_issue(
                    "编码格式错误",
                    "基础",
                    [row_no],
                    f"{column} 应为 {width} 位数字，当前值为 {value}",
                    city=record.get("地市", ""),
                    city_code=record.get("地市代码", ""),
                    department=record.get("主管部门", ""),
                    department_code=record.get("主管部门编码", ""),
                    unit=record.get("报考单位", ""),
                    unit_code=record.get("报考单位编码", ""),
                    job=record.get("报考岗位", ""),
                    job_code=record.get("报考岗位编码", ""),
                )

    detect_mapping_conflicts(records, scope_keys=(), code_key="地市代码", name_key="地市", level="地市")
    detect_mapping_conflicts(records, scope_keys=("地市代码",), code_key="主管部门编码", name_key="主管部门", level="主管部门")
    detect_mapping_conflicts(
        records,
        scope_keys=("地市代码", "主管部门编码"),
        code_key="报考单位编码",
        name_key="报考单位",
        level="报考单位",
    )
    detect_mapping_conflicts(
        records,
        scope_keys=("地市代码", "主管部门编码", "报考单位编码"),
        code_key="报考岗位编码",
        name_key="报考岗位",
        level="报考岗位",
    )

    detect_sequence_issues(
        records,
        parent_keys=("地市代码",),
        child_code_key="主管部门编码",
        level="主管部门",
        width=3,
    )
    detect_sequence_issues(
        records,
        parent_keys=("地市代码", "主管部门编码"),
        child_code_key="报考单位编码",
        level="报考单位",
        width=3,
    )
    detect_sequence_issues(
        records,
        parent_keys=("地市代码", "主管部门编码", "报考单位编码"),
        child_code_key="报考岗位编码",
        level="报考岗位",
        width=2,
    )

    unique_issues: list[AuditIssue] = []
    seen: set[tuple[str, ...]] = set()
    for issue in issues:
        key = (
            issue.issue_type,
            issue.level,
            issue.excel_rows,
            issue.description,
            issue.city,
            issue.city_code,
            issue.department,
            issue.department_code,
            issue.unit,
            issue.unit_code,
            issue.job,
            issue.job_code,
        )
        if key in seen:
            continue
        seen.add(key)
        unique_issues.append(issue)

    return AuditResult(source_path=excel_path, total_rows=len(records), issues=unique_issues)


def _issue_sort_key(issue: AuditIssue) -> tuple[object, ...]:
    def _row_key() -> int:
        first = (issue.excel_rows or "").split("、", 1)[0].strip()
        return int(first) if first.isdigit() else 10**9

    return (
        issue.city_code,
        issue.department_code,
        issue.unit_code,
        issue.job_code,
        LEVEL_ORDER.get(issue.level, 99),
        issue.issue_type,
        _row_key(),
    )


def _plain_description(issue: AuditIssue) -> str:
    if issue.issue_type == "缺少必填值":
        return issue.description
    if issue.issue_type == "编码格式错误":
        return issue.description
    if issue.issue_type == "编码映射冲突":
        return f"同一个编码对应了多个名称。{issue.description}"
    if issue.issue_type == "名称映射冲突":
        return f"同一个名称对应了多个编码。{issue.description}"
    if issue.issue_type == "顺延起始错误":
        return f"当前层级编码不是从起始值开始。{issue.description}"
    return issue.description


def _level_path(issue: AuditIssue) -> str:
    parts: list[str] = []
    if issue.city:
        parts.append(issue.city)
    if issue.department:
        parts.append(issue.department)
    elif issue.level in {"主管部门", "报考单位", "报考岗位"} and issue.department_code:
        parts.append(f"主管部门[{issue.department_code}]")
    if issue.unit:
        parts.append(issue.unit)
    elif issue.level in {"报考单位", "报考岗位"} and issue.unit_code:
        parts.append(f"报考单位[{issue.unit_code}]")
    if issue.job:
        parts.append(issue.job)
    elif issue.level == "报考岗位" and issue.job_code:
        parts.append(f"报考岗位[{issue.job_code}]")
    return " / ".join(parts)


def _sheet_headers() -> list[str]:
    return [
        "地市",
        "地市代码",
        "主管部门",
        "主管部门编码",
        "报考单位",
        "报考单位编码",
        "报考岗位",
        "报考岗位编码",
        "层级路径",
        "问题层级",
        "问题类型",
        "问题说明",
    ]


def _format_issue_sheet(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
    for column in "ABCDEFGHIJKL":
        sheet.column_dimensions[column].width = 18
    sheet.column_dimensions["I"].width = 36
    sheet.column_dimensions["K"].width = 18
    sheet.column_dimensions["L"].width = 56


def _write_issue_row(sheet, issue: AuditIssue) -> None:
    sheet.append(
        [
            issue.city,
            issue.city_code,
            issue.department,
            issue.department_code,
            issue.unit,
            issue.unit_code,
            issue.job,
            issue.job_code,
            _level_path(issue),
            issue.level,
            issue.issue_type,
            _plain_description(issue),
        ]
    )
    issue_type_cell = sheet.cell(sheet.max_row, 11)
    issue_type_cell.fill = ISSUE_FILL
    issue_type_cell.font = ISSUE_FONT


def _safe_file_stem(raw: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid else char for char in raw).strip(" .")
    return cleaned or "未填写地市"


def export_job_audit_result(result: AuditResult, output_dir: Path) -> list[Path]:
    headers = _sheet_headers()
    city_groups: dict[tuple[str, str], list[AuditIssue]] = defaultdict(list)
    for issue in sorted(result.issues, key=_issue_sort_key):
        city_groups[(issue.city or "未填写地市", issue.city_code or "")].append(issue)

    output_dir.mkdir(parents=True, exist_ok=True)
    exported_paths: list[Path] = []

    master_workbook = Workbook()
    master_summary = master_workbook.active
    master_summary.title = "校验摘要"
    master_summary.append(["源文件", str(result.source_path)])
    master_summary.append(["数据总行数", result.total_rows])
    master_summary.append(["问题总条数", result.issue_count])
    master_summary.append(["地市文件数", len(city_groups)])
    master_summary.append([])
    master_summary.append(["说明", "总表保留校验摘要；分地市文件仅保留问题明细，且已去掉 Excel 行号。"])

    master_issue_sheet = master_workbook.create_sheet("问题明细")
    master_issue_sheet.append(headers)
    for issue in sorted(result.issues, key=_issue_sort_key):
        _write_issue_row(master_issue_sheet, issue)
    _format_issue_sheet(master_issue_sheet)

    master_path = output_dir / f"{_safe_file_stem(result.source_path.stem)}_岗位核对总表.xlsx"
    master_workbook.save(master_path)
    master_workbook.close()
    exported_paths.append(master_path)

    for city_index, ((city, city_code), issues) in enumerate(city_groups.items(), start=1):
        workbook = Workbook()
        issue_sheet = workbook.active
        issue_sheet.title = "问题明细"
        issue_sheet.append(headers)
        for issue in issues:
            _write_issue_row(issue_sheet, issue)
        _format_issue_sheet(issue_sheet)

        file_stem = _safe_file_stem(f"{city_code}_{city}" if city_code else city)
        output_path = output_dir / f"{file_stem}_岗位核对结果.xlsx"
        if output_path.exists():
            output_path = output_dir / f"{file_stem}_{city_index}_岗位核对结果.xlsx"
        workbook.save(output_path)
        workbook.close()
        exported_paths.append(output_path)

    return exported_paths
